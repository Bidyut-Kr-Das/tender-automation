import asyncio
import logging
from pathlib import Path

from playwright.sync_api import sync_playwright
from django.conf import settings

from automation_v2.lib.browser import detect_chrome_path
from automation_v2.lib.result_scrapers import (
    _click_indian_result,
    _click_today_results,
    _download_today_excel,
    _iter_rows as _iter_247_rows,
    _map_stage,
    _tender247_login,
)

if hasattr(asyncio, "WindowsProactorEventLoopPolicy"):
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

logger = logging.getLogger(__name__)


def _parse_247_excel(path: Path) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return []
    headers: list[str] = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(c).strip() if c is not None else "" for c in row]
        break
    headers = [h for h in headers if h]
    wb.close()
    logger.info("[Tender247Result] Excel headers: %s", headers)
    print(f"[Tender247Result] Excel headers: {headers}")
    return headers


def _log_247_updates(path: Path) -> list[dict]:
    from .tender_result_updater import update_tender_result

    gen = _iter_247_rows(path)
    try:
        headers_raw = next(gen)
    except StopIteration:
        return []
    headers = [str(c).strip() if c is not None else "" for c in headers_raw]
    hmap = {h.lower(): idx for idx, h in enumerate(headers)}
    idx_ref = hmap.get("tender reference no")
    idx_l1 = hmap.get("winner bidder")
    idx_amt = hmap.get("contract value")
    idx_comp = hmap.get("participator bidders")
    idx_status = hmap.get("tender stage")
    if idx_ref is None or idx_l1 is None or idx_amt is None:
        logger.warning("[Tender247Result] missing columns headers=%s", headers)
        print(f"[Tender247Result] missing columns headers={headers}")
        return []
    updates: list[dict] = []
    for r in gen:
        if not r:
            continue
        ref = str(r[idx_ref]).strip() if r[idx_ref] is not None else ""
        if not ref:
            continue
        l1 = str(r[idx_l1]).strip() if r[idx_l1] is not None else ""
        amt_raw = str(r[idx_amt]).strip() if r[idx_amt] is not None else ""
        competitors = str(r[idx_comp]).strip() if idx_comp is not None and r[idx_comp] is not None else ""
        cur_raw = str(r[idx_status]).strip() if idx_status is not None and r[idx_status] is not None else ""
        cur_status = _map_stage(cur_raw)
        res = update_tender_result(ref, l1, amt_raw, competitors or None, cur_status or None)
        updates.append(res)
    logger.info("[Tender247Result] total rows %d", len(updates))
    print(f"[Tender247Result] total rows {len(updates)}")
    return updates


def get_tender247_result(email: str | None = None, password: str | None = None) -> dict:
    email = email or settings.TENDER247_EMAIL
    password = password or settings.TENDER247_PASSWORD
    if not email or not password:
        return {"success": False, "error": "TENDER247_EMAIL/PASSWORD not configured"}

    chrome_path = detect_chrome_path()
    tenders_file: Path | None = None
    url = ""
    # Phase A: Playwright
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=chrome_path, headless=settings.HEADLESS_BROWSER)
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        try:
            ok = _tender247_login(page, email, password)
            if not ok:
                return {"success": False, "error": "Tender247 login failed", "url": page.url}
            _click_indian_result(page)
            _click_today_results(page)
            tenders_file = _download_today_excel(page)
            url = page.url
        except Exception as e:
            logger.exception("[Tender247Result] failed (download): %s", e)
            return {"success": False, "error": str(e), "url": page.url if 'page' in locals() else ""}
        finally:
            try:
                browser.close()
            except Exception:
                pass
    # clear async loop before DB like tiger
    try:
        asyncio.set_event_loop(None)
    except Exception:
        pass
    try:
        from django.db import close_old_connections
        close_old_connections()
    except Exception:
        pass
    # Phase B: DB
    try:
        assert tenders_file is not None
        headers = _parse_247_excel(tenders_file)
        updates = _log_247_updates(tenders_file)
        return {"success": True, "url": url, "tenders_file": str(tenders_file), "headers": headers, "updates": updates, "would_update": sum(1 for u in updates if u["found"])}
    except Exception as e:
        logger.exception("[Tender247Result] failed (DB): %s", e)
        return {"success": False, "error": str(e)}
