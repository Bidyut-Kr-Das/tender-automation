import asyncio
import logging
from pathlib import Path

from playwright.sync_api import sync_playwright
from django.conf import settings

from automation_v2.lib.browser import detect_chrome_path
from automation_v2.lib.result_scrapers import _click_result_tab, _click_tenders_download, _iter_rows as _iter_tiger_rows
from automation_v2.lib.tender_tiger import _tiger_login

if hasattr(asyncio, "WindowsProactorEventLoopPolicy"):
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

logger = logging.getLogger(__name__)


def _parse_tenders_excel(path: Path) -> list[str]:
    # ponytail: read_only first row only — full parse when spec lands
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return []
    headers: list[str] = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(c).strip() if c is not None else "" for c in row]
        break
    headers = [h for h in headers if h]
    wb.close()
    logger.info("[TigerResult] Excel headers: %s", headers)
    print(f"[TigerResult] Excel headers: {headers}")
    return headers


def _log_tenders_updates(path: Path) -> list[dict]:
    # ponytail: delegates to reusable updater — no DB logic here
    from .tender_result_updater import update_tender_result

    gen = _iter_tiger_rows(path)
    try:
        headers_raw = next(gen)
    except StopIteration:
        return []
    headers = [str(c).strip() if c is not None else "" for c in headers_raw]
    hmap = {h.lower(): idx for idx, h in enumerate(headers)}
    idx_ref = hmap.get("actual tender number")
    idx_l1 = hmap.get("l1")
    idx_amt = hmap.get("contract amount")
    if idx_ref is None or idx_l1 is None or idx_amt is None:
        logger.warning("[TigerResult] missing columns headers=%s", headers)
        print(f"[TigerResult] missing columns headers={headers}")
        return []
    updates: list[dict] = []
    for r in gen:
        if not r:
            continue
        ref = str(r[idx_ref]).strip() if r[idx_ref] is not None else ""
        if not ref:
            continue
        l1 = str(r[idx_l1]).strip() if r[idx_l1] is not None else ""
        amt_raw = str(r[idx_amt]).strip() if r[idx_amt] is not None else ""
        res = update_tender_result(ref, l1, amt_raw)
        updates.append(res)
    logger.info("[TigerResult] total rows %d", len(updates))
    print(f"[TigerResult] total rows {len(updates)}")
    return updates


def get_tender_tiger_result(email: str | None = None, password: str | None = None) -> dict:
    email = email or settings.TENDER_TIGER_EMAIL
    password = password or settings.TENDER_TIGER_PASSWORD
    if not email or not password:
        return {"success": False, "error": "TENDER_TIGER_EMAIL/PASSWORD not configured"}

    chrome_path = detect_chrome_path()
    tenders_file: Path | None = None
    url = ""
    # Phase A: Playwright (async loop)
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=chrome_path, headless=settings.HEADLESS_BROWSER)
        page = browser.new_page()
        try:
            ok = _tiger_login(page, email, password)
            if not ok:
                return {"success": False, "error": "Tiger login failed", "url": page.url}
            page.wait_for_timeout(2000)
            _click_result_tab(page)
            tenders_file = _click_tenders_download(page)
            url = page.url
        except Exception as e:
            logger.exception("[TigerResult] failed (download): %s", e)
            return {"success": False, "error": str(e)}
        finally:
            browser.close()
    # clear async loop before DB (sync) — ponytail: Proactor loop from Playwright confuses Django ORM
    try:
        asyncio.set_event_loop(None)
    except Exception:
        pass
    try:
        from django.db import close_old_connections

        close_old_connections()
    except Exception:
        pass
    # Phase B: DB (sync, no async loop)
    try:
        assert tenders_file is not None
        headers = _parse_tenders_excel(tenders_file)
        updates = _log_tenders_updates(tenders_file)
        return {"success": True, "url": url, "tenders_file": str(tenders_file), "headers": headers, "updates": updates, "would_update": sum(1 for u in updates if u["found"])}
    except Exception as e:
        logger.exception("[TigerResult] failed (DB): %s", e)
        return {"success": False, "error": str(e)}
