"""TenderTiger / Tender247 result Excel scrapers. No DB code; used by both automation_v2 and tender_search."""
import logging
import re
from pathlib import Path

from django.conf import settings
from playwright.sync_api import sync_playwright

from .browser import detect_chrome_path
from .tender_tiger import _tiger_login

logger = logging.getLogger(__name__)


# --- TenderTiger ---

def _click_result_tab(page) -> None:
    # ponytail: exact locator — no generic fallback
    loc = page.locator("a.tab-item.aiSearchTab.airesult-tab[title='Result']")
    loc.first.wait_for(state="visible", timeout=15000)
    loc.first.scroll_into_view_if_needed(timeout=5000)
    print(f"[TigerResult] Clicking Result tab")
    loc.first.click()
    page.wait_for_timeout(2000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass


def _click_tenders_download(page) -> Path:
    # ponytail: parent div.filter-btn-listing-tr → span — scopes to Tenders button
    parent = page.locator('div.filter-btn-listing-tr')
    parent.first.wait_for(state="visible", timeout=15000)
    loc = parent.locator('span.icon-label.arrow-none')
    loc.first.wait_for(state="visible", timeout=15000)
    loc.first.scroll_into_view_if_needed(timeout=5000)
    print(f"[TigerResult] Clicking span.icon-label.arrow-none")
    with page.expect_download(timeout=60000) as dl:
        try:
            loc.first.click(force=True)
        except Exception:
            loc.first.evaluate("el => el.click()")
    download = dl.value
    download_dir = Path(settings.TENDER_PARSING_TEMP_DIR)
    download_dir.mkdir(parents=True, exist_ok=True)
    suggested = download.suggested_filename or "tenders.xlsx"
    dest = download_dir / suggested
    download.save_as(str(dest))
    print(f"[TigerResult] Downloaded: {dest}")
    if not dest.exists() or dest.stat().st_size == 0:
        raise RuntimeError(f"Tenders excel not captured: {dest}")
    return dest


# --- Tender247 ---

def _tender247_login(page, email: str, password: str) -> bool:
    print("[Tender247Result] Navigating to https://www.tender247.com/auth/tender ...")
    page.goto("https://www.tender247.com/auth/tender", timeout=50000)
    try:
        page.locator("input[name='emailId']").first.wait_for(state="visible", timeout=3000)
    except Exception:
        login_btn = page.get_by_role("button", name="Log in")
        if login_btn.count() == 0:
            login_btn = page.locator("button:has-text('Log in')")
        if login_btn.count() > 0:
            try:
                if login_btn.first.is_visible():
                    print("[Tender247Result] Clicking Log in...")
                    login_btn.first.click()
                    page.wait_for_timeout(1500)
            except Exception:
                pass
    signup_btn = page.locator("button:has-text('Sign Up')")
    if signup_btn.count() > 0:
        try:
            if signup_btn.first.is_visible():
                signup_btn.first.click()
                page.wait_for_timeout(1000)
        except Exception:
            pass
    page.locator("input[name='emailId']").first.wait_for(state="visible", timeout=10000)
    page.locator("input[name='emailId']").first.fill(email)
    page.locator("input[name='password']").first.fill(password)
    page.locator("button[type='submit']:has-text('Submit')").click()
    page.wait_for_timeout(5000)
    close_btn = page.locator("button:has(span.sr-only:text('Close'))")
    if close_btn.count() > 0:
        try:
            if close_btn.first.is_visible():
                print("[Tender247Result] Closing dialog...")
                close_btn.first.click()
                page.wait_for_timeout(1000)
        except Exception:
            pass
    # ponytail: same success heuristic as non_gem_tender_pdf_downloader login_tender247
    try:
        signup_text = page.locator("button:has-text('Sign Up')").inner_text() if page.locator("button:has-text('Sign Up')").count() > 0 else ""
    except Exception:
        signup_text = ""
    return "Sign Up" not in signup_text


def _click_indian_result(page) -> None:
    print("[Tender247Result] Clicking Result menu...")
    result_loc = page.locator("li.cursor-pointer:has(img[src*='tender-result-icon']):has-text('Result')")
    result_loc.first.wait_for(state="visible", timeout=15000)
    result_loc.first.scroll_into_view_if_needed(timeout=5000)
    result_loc.first.click()
    page.wait_for_timeout(1500)
    try:
        page.wait_for_load_state("networkidle", timeout=10000)
    except Exception:
        pass
    print("[Tender247Result] Clicking Indian...")
    # ponytail: exact href scopes away from global Indian text dupes
    candidates = [
        "a[href='/auth/result']:has-text('Indian')",
        "a[href='/auth/result']",
        "a:has(img[src*='domestic-icon']):has-text('Indian')",
        "a:has-text('Indian')",
    ]
    clicked = False
    for sel in candidates:
        loc = page.locator(sel).first
        try:
            if loc.count() > 0 and loc.is_visible():
                print(f"[Tender247Result] Indian via {sel!r}")
                loc.click()
                clicked = True
                break
        except Exception:
            continue
    if not clicked:
        loc = page.get_by_text("Indian", exact=True).first
        loc.wait_for(state="visible", timeout=15000)
        loc.click()
    page.wait_for_timeout(2000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass


def _click_today_results(page) -> None:
    print("[Tender247Result] Clicking Today Results...")
    # ponytail: card has img alt Today Results + p text — scopes away from Today Tenders
    loc = page.locator("div.cursor-pointer:has(p:has-text('Today Results'))")
    if loc.count() == 0:
        loc = page.locator("div.cursor-pointer:has(img[alt='Today Results'])")
    loc.first.wait_for(state="visible", timeout=15000)
    loc.first.scroll_into_view_if_needed(timeout=5000)
    loc.first.click()
    page.wait_for_timeout(2000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass


def _download_today_excel(page) -> Path:
    print("[Tender247Result] Waiting Today Results data...")
    page.wait_for_timeout(2000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass
    print("[Tender247Result] Clicking Download Excel...")
    loc = page.locator("span:has(img[alt='Download Excel'])")
    if loc.count() == 0:
        loc = page.locator("img[alt='Download Excel']")
    loc.first.wait_for(state="visible", timeout=15000)
    loc.first.scroll_into_view_if_needed(timeout=5000)
    download_dir = Path(settings.TENDER_PARSING_TEMP_DIR)
    download_dir.mkdir(parents=True, exist_ok=True)
    with page.expect_download(timeout=30000) as dl:
        try:
            loc.first.click()
        except Exception:
            loc.first.evaluate("el => el.click()")
    download = dl.value
    suggested = download.suggested_filename or "tender247_today.xlsx"
    dest = download_dir / suggested
    download.save_as(str(dest))
    print(f"[Tender247Result] Downloaded: {dest}")
    if not dest.exists() or dest.stat().st_size == 0:
        raise RuntimeError(f"Tender247 excel not captured: {dest}")
    return dest


def _map_stage(raw: str) -> str:
    # ponytail: substring match on Tender Stage → normalized status
    low = (raw or "").strip().lower()
    if re.search(r"\baoc\b", low):
        return "AWARDED"
    if "financial" in low:
        return "FINANCIAL EVALUATION"
    if "technical" in low:
        return "TECHNICAL BID OPENED"
    return ""


# --- Excel rows ---

def _iter_rows(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return
    try:
        for row in ws.iter_rows(values_only=True):
            yield row
    finally:
        wb.close()


def _cr_to_number(val: str | None) -> str:
    if not val:
        return ""
    s = str(val).strip()
    m = re.search(r"([\d,]+(?:\.\d+)?)\s*cr", s, re.I)
    if m:
        num = m.group(1).replace(",", "")
        try:
            return str(int(float(num) * 1e7))
        except ValueError:
            return num
    return re.sub(r"[^0-9.]", "", s)


TIGER_COLUMNS = {"ref": "actual tender number", "l1": "l1", "amount": "contract amount"}
T247_COLUMNS = {"ref": "tender reference no", "l1": "winner bidder", "amount": "contract value",
                "competitors": "participator bidders", "stage": "tender stage"}


def read_rows(path: Path, columns: dict) -> list[dict]:
    """Result rows as plain dicts. `columns` maps ref/l1/amount (required) and competitors/stage (optional)
    to lowercase header names. Raises ValueError if a required column is missing."""
    gen = _iter_rows(path)
    try:
        headers_raw = next(gen)
    except StopIteration:
        return []
    hmap = {str(c).strip().lower(): i for i, c in enumerate(headers_raw) if c is not None}
    idx = {k: hmap.get(h) for k, h in columns.items()}
    if any(idx.get(k) is None for k in ("ref", "l1", "amount")):
        raise ValueError(f"missing columns, headers={[h for h in hmap]}")

    def cell(r, key) -> str:
        i = idx.get(key)
        return str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else ""

    rows = []
    for r in gen:
        if not r or not cell(r, "ref"):
            continue
        parts = re.split(r"<br\s*/?>", cell(r, "ref"), flags=re.I)
        l1, amount, stage = cell(r, "l1"), cell(r, "amount"), cell(r, "stage")
        rows.append({
            "referenceNo": parts[0].strip(),
            "reverseAuction": len(parts) > 1 and parts[1].strip() != "",
            "l1": l1,
            "isLaser": "laser power" in l1.lower(),
            "contractAmount": amount,
            "contractValue": _cr_to_number(amount),
            "competitors": cell(r, "competitors") or None,
            "tenderStage": stage or None,
            "currentStatus": _map_stage(stage) or None,
        })
    return rows


# --- Download (Playwright) ---

def _download(login, steps, email: str, password: str, name: str, **page_opts) -> Path:
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=detect_chrome_path(), headless=settings.HEADLESS_BROWSER)
        try:
            page = browser.new_page(**page_opts)
            if not login(page, email, password):
                raise RuntimeError(f"{name} login failed (url={page.url})")
            *clicks, download = steps
            for click in clicks:
                click(page)
            return download(page)
        finally:
            browser.close()


def download_tiger_results(email: str, password: str) -> Path:
    def _wait(page):
        page.wait_for_timeout(2000)
    return _download(_tiger_login, (_wait, _click_result_tab, _click_tenders_download), email, password, "Tiger")


def download_247_results(email: str, password: str) -> Path:
    return _download(_tender247_login, (_click_indian_result, _click_today_results, _download_today_excel),
                     email, password, "Tender247", viewport={"width": 1920, "height": 1080})
